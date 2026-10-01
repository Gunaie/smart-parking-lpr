# -*- coding: utf-8 -*-
"""命令行版停车流程：集成车牌识别、Excel 记录和 MQTT 闸机控制。

进场: python test_parking.py --entry <image>
离场: python test_parking.py --exit  <image>
四种场景：正常进场 / 重复进场(REJECT) / 正常离场 / 无记录离场(NO RECORD)
"""
import argparse
import time
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

from config import settings
from plate_pipeline import PlateRecognizer
from gate_control import open_gate, close_gate

EXCEL_PATH = settings.EXCEL_PATH
HEADERS = ["License Plate", "Entry Time"]
GATE_DELAY = settings.GATE_DELAY


def _load_sheet():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb["In Lot"]
    return wb, ws


def _ensure_excel():
    if not Path(EXCEL_PATH).exists():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "In Lot"
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        wb.save(EXCEL_PATH)
        print(f"[Excel] 创建 {EXCEL_PATH}")


def _exists(plate: str) -> bool:
    _, ws = _load_sheet()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] == plate:
            return True
    return False


def _add(plate: str):
    wb, ws = _load_sheet()
    ws.append([plate, datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    wb.save(EXCEL_PATH)


def _delete(plate: str):
    wb, ws = _load_sheet()
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row and row[0] == plate:
            ws.delete_rows(i)
            wb.save(EXCEL_PATH)
            return


def _recognize(image_path, model_path=None) -> str | None:
    recognizer = PlateRecognizer(model_path)
    results = recognizer(image_path, save=True)
    if not results:
        return None
    return max(results, key=lambda r: r["confidence"])["plate"]


def do_entry(image_path: str, model_path: str = None):
    print(f"[MODE] Entry  image={image_path}")
    plate = _recognize(image_path, model_path)
    if not plate:
        print("-> Entry failed: no plate detected")
        return
    if _exists(plate):
        print(f"[REJECT] {plate} already in lot")
        return
    _add(plate)
    print(f"[ENTRY] {plate}  [{datetime.now():%Y-%m-%d %H:%M:%S}]")
    open_gate()
    time.sleep(GATE_DELAY)
    close_gate()


def do_exit(image_path: str, model_path: str = None):
    print(f"[MODE] Exit  image={image_path}")
    plate = _recognize(image_path, model_path)
    if not plate:
        print("-> Exit failed: no plate detected")
        return
    if not _exists(plate):
        print(f"[NO RECORD] {plate}")
        return
    _delete(plate)
    print(f"[EXIT] {plate}")
    open_gate()
    time.sleep(GATE_DELAY)
    close_gate()


def main():
    parser = argparse.ArgumentParser(description="命令行停车流程测试")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--entry", metavar="IMAGE", help="进场识别")
    group.add_argument("--exit", metavar="IMAGE", help="离场识别")
    parser.add_argument("--model", default=None, help="YOLO 权重路径")
    args = parser.parse_args()

    settings.ensure_runtime_dirs()
    _ensure_excel()

    if args.entry:
        do_entry(args.entry, args.model)
    else:
        do_exit(args.exit, args.model)


if __name__ == "__main__":
    main()
