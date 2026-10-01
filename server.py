# -*- coding: utf-8 -*-
"""FastAPI 智能社区车牌识别综合服务。

启动：
    python server.py                   # 完整链路
    python server.py --no-gate         # 跳过闸机（无 MQTT/舵机时用）
    python server.py --no-tts          # 跳过语音
    python server.py --port 8000
"""
import argparse
import asyncio
import io
import json
import os
import time
from datetime import datetime
from pathlib import Path

import cv2
import openpyxl
import uvicorn
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse
from openpyxl.styles import Alignment, Font
from sse_starlette.sse import EventSourceResponse

from config import settings
from plate_pipeline import PlateRecognizer
import tts_utils

EXCEL_HEADERS = ["License Plate", "Entry Time", "Space"]

# ---------- 全局状态 ----------
app = FastAPI(title="智能社区车牌识别系统")
recognizer: PlateRecognizer | None = None
sse_queues: set[asyncio.Queue] = set()
camera: cv2.VideoCapture | None = None
gate_enabled: bool = True
tts_enabled: bool = True


# ===================== Excel 管理 =====================
def _load_sheet():
    if not Path(settings.EXCEL_PATH).exists():
        _ensure_excel()
    wb = openpyxl.load_workbook(settings.EXCEL_PATH)
    if "In Lot" not in wb.sheetnames:
        ws = wb.create_sheet("In Lot")
    else:
        ws = wb["In Lot"]
    # 兼容旧两列表：补第三列表头
    headers = [ws.cell(1, c).value for c in range(1, 4)]
    if headers != EXCEL_HEADERS:
        for c, h in enumerate(EXCEL_HEADERS, start=1):
            cell = ws.cell(1, c, h)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        wb.save(settings.EXCEL_PATH)
    return wb, ws


def _ensure_excel():
    settings.RECORD_DIR.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "In Lot"
    ws.append(EXCEL_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    wb.save(settings.EXCEL_PATH)


def _read_records() -> list[dict]:
    _, ws = _load_sheet()
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        records.append({
            "plate": row[0],
            "entry_time": row[1] if len(row) > 1 else "",
            "space": row[2] if len(row) > 2 else "",
        })
    return records


def _exists(plate: str) -> dict | None:
    for r in _read_records():
        if r["plate"] == plate:
            return r
    return None


def _add(plate: str, space: str):
    wb, ws = _load_sheet()
    ws.append([plate, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), space])
    wb.save(settings.EXCEL_PATH)


def _delete(plate: str):
    wb, ws = _load_sheet()
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row and row[0] == plate:
            ws.delete_rows(i)
            wb.save(settings.EXCEL_PATH)
            return
    wb.close()


def _allocated_spaces() -> set[str]:
    return {r["space"] for r in _read_records() if r.get("space")}


def _free_space() -> str | None:
    used = _allocated_spaces()
    for s in settings.ALL_SPACES:
        if s not in used:
            return s
    return None


# ===================== 闸机 / TTS =====================
def _gate_open_close():
    if not gate_enabled:
        return
    try:
        from gate_control import open_gate, close_gate

        open_gate()
        time.sleep(settings.GATE_DELAY)
        close_gate()
    except Exception as e:
        print(f"[Gate] 异常: {e}")


def _speak(text):
    if not tts_enabled:
        return
    try:
        tts_utils.speak(text)
    except Exception as e:
        print(f"[TTS] 异常: {e}")


# ===================== SSE =====================
async def sse_broadcast(event: dict):
    if not sse_queues:
        return
    for q in list(sse_queues):
        try:
            await q.put(event)
        except Exception:
            pass


async def sse_generator(request: Request):
    q: asyncio.Queue = asyncio.Queue()
    sse_queues.add(q)
    try:
        # 先推一次当前状态（data 必须是字符串，前端 JSON.parse）
        yield {"event": "state", "data": json.dumps(state_payload())}
        while True:
            if await request.is_disconnected():
                break
            try:
                event = await asyncio.wait_for(q.get(), timeout=15)
                yield event
            except asyncio.TimeoutError:
                yield {"event": "ping", "data": "pong"}
    finally:
        sse_queues.discard(q)


# ===================== 状态 =====================
def state_payload() -> dict:
    records = _read_records()
    spaces = {}
    for r in records:
        if r.get("space"):
            spaces[r["space"]] = r
    occupied = len(spaces)
    return {
        "total": settings.TOTAL_SPACES,
        "occupied": occupied,
        "free": settings.TOTAL_SPACES - occupied,
        "spaces": spaces,
        "records": records,
    }


@app.get("/state")
async def state():
    return state_payload()


# ===================== 业务接口 =====================
@app.post("/entry")
async def entry(file: UploadFile = File(...)):
    img_bytes = await file.read()
    plate, det = await asyncio.to_thread(_recognize_one, img_bytes)
    if not plate:
        raise HTTPException(status_code=422, detail="未检测到车牌")

    if _exists(plate):
        raise HTTPException(status_code=409, detail="重复进场")

    space = _free_space()
    if not space:
        raise HTTPException(status_code=409, detail="车位已满")

    _add(plate, space)
    payload = {"type": "entry", "plate": plate, "space": space,
               "entry_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
    await sse_broadcast({"event": "entry", "data": json.dumps(payload)})

    # 闸机和语音放后台线程，避免阻塞响应
    asyncio.create_task(asyncio.to_thread(_gate_open_close))
    asyncio.create_task(asyncio.to_thread(
        _speak, f"{plate}，欢迎光临，{space}号车位"))

    return {"status": "ok", **payload}


@app.post("/exit")
async def exit(file: UploadFile = File(...)):
    img_bytes = await file.read()
    plate, det = await asyncio.to_thread(_recognize_one, img_bytes)
    if not plate:
        raise HTTPException(status_code=422, detail="未检测到车牌")

    rec = _exists(plate)
    if not rec:
        raise HTTPException(status_code=404, detail="无入场记录")

    space = rec.get("space", "")
    _delete(plate)
    payload = {"type": "exit", "plate": plate, "space": space}
    await sse_broadcast({"event": "exit", "data": json.dumps(payload)})

    asyncio.create_task(asyncio.to_thread(_gate_open_close))
    asyncio.create_task(asyncio.to_thread(_speak, f"{plate}，一路顺风"))
    return {"status": "ok", **payload}


@app.post("/reset")
async def reset():
    wb, ws = _load_sheet()
    ws.delete_rows(2, ws.max_row)
    wb.save(settings.EXCEL_PATH)
    await sse_broadcast({"event": "reset", "data": "{}"})
    return {"status": "ok"}


# ===================== 摄像头 / MJPEG =====================
def _get_camera() -> cv2.VideoCapture | None:
    """Linux/Docker 容器中通常无摄像头，返回 None 表示不可用。"""
    global camera
    if camera is None:
        if os.name != "nt":
            return None  # 容器/服务器环境无摄像头
        camera = cv2.VideoCapture(settings.CAMERA_INDEX, cv2.CAP_DSHOW)
        if not camera.isOpened():
            camera = None
            return None
    elif not camera.isOpened():
        camera.release()
        camera = None
        if os.name != "nt":
            return None
        camera = cv2.VideoCapture(settings.CAMERA_INDEX, cv2.CAP_DSHOW)
        if not camera.isOpened():
            camera = None
            return None
    return camera


def _jpeg_frame():
    cap = _get_camera()
    if cap is None:
        return None
    ok, frame = cap.read()
    if not ok:
        return None
    _, buf = cv2.imencode(".jpg", frame)
    return buf.tobytes()


@app.get("/snapshot")
async def snapshot():
    jpg = await asyncio.to_thread(_jpeg_frame)
    if not jpg:
        raise HTTPException(status_code=500, detail="摄像头读取失败")
    return StreamingResponse(io.BytesIO(jpg), media_type="image/jpeg")


@app.get("/video")
async def video():
    def gen():
        boundary = "--frame"
        while True:
            jpg = _jpeg_frame()
            if not jpg:
                time.sleep(0.1)
                continue
            yield (f"{boundary}\r\nContent-Type: image/jpeg\r\n"
                   f"Content-Length: {len(jpg)}\r\n\r\n").encode() + jpg + b"\r\n"
            time.sleep(0.05)

    return StreamingResponse(gen(), media_type="multipart/x-mixed-replace; boundary=frame")


@app.get("/stream")
async def stream(request: Request):
    return EventSourceResponse(sse_generator(request))


# ===================== 识别辅助 =====================
def _recognize_one(img_bytes: bytes) -> tuple[str | None, dict | None]:
    import numpy as np
    import tempfile

    arr = np.frombuffer(img_bytes, np.uint8)
    img = cv2.imdecode(arr, cv2.IMREAD_COLOR)
    if img is None:
        return None, None
    # PlateRecognizer.__call__ 接受文件路径，写临时文件
    with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False, dir=str(settings.UPLOADS_DIR)) as tf:
        tf.write(img_bytes)
        tmp_path = tf.name
    try:
        results = recognizer(tmp_path)
    finally:
        try:
            Path(tmp_path).unlink()
        except Exception:
            pass
    if not results:
        return None, None
    best = max(results, key=lambda r: r["confidence"])
    return best["plate"], best


# ===================== 页面 =====================
@app.get("/", response_class=HTMLResponse)
async def index():
    return (Path(__file__).parent / "static" / "index.html").read_text(
        encoding="utf-8")


# ===================== 启动 =====================
def main():
    global recognizer, gate_enabled, tts_enabled
    parser = argparse.ArgumentParser(description="智能社区车牌识别服务")
    parser.add_argument("--host", default=settings.HOST)
    parser.add_argument("--port", type=int, default=settings.PORT)
    parser.add_argument("--no-gate", action="store_true", help="跳过闸机")
    parser.add_argument("--no-tts", action="store_true", help="跳过语音")
    parser.add_argument("--model", default=None, help="YOLO 权重路径")
    args = parser.parse_args()

    gate_enabled = not args.no_gate
    tts_enabled = not args.no_tts
    settings.ensure_runtime_dirs()
    _ensure_excel()

    print("[Init] 加载识别流水线...")
    recognizer = PlateRecognizer(args.model)
    print(f"[Init] gate={gate_enabled} tts={tts_enabled}")
    print(f"[Init] http://{args.host}:{args.port}")

    uvicorn.run(app, host=args.host, port=args.port, log_level="info")


if __name__ == "__main__":
    main()
